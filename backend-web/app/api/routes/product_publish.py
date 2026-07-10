"""
商品发布 API 路由

功能：
1. 素材库管理（CRUD）
2. 单品发布（触发 Playwright 自动化）
3. 批量发布（后台任务异步执行）
4. 发布日志查询（分页+过滤）
"""
from __future__ import annotations

from typing import Any, Dict, List, Optional

import hashlib
import io
import os
import re
import time
import uuid
import zipfile
from pathlib import Path
from urllib.parse import unquote, urlparse

from fastapi import APIRouter, BackgroundTasks, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from loguru import logger
from openpyxl import Workbook, load_workbook
from pydantic import BaseModel, Field
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import text, update, select

from app.api.deps import get_current_active_user, get_db_session
from app.services.product_publish_service import ProductMaterialService
from app.services.publish_batch_status_service import PublishBatchStatusService
from app.services.publish_execution_service import PublishExecutorService, PublishLogService
from common.models.user import User, UserRole
from common.models.product_material import ProductMaterial
from common.schemas.common import ApiResponse
from common.utils.local_image_upload import (
    ImageUploadError,
    read_image_with_size_check,
    safe_image_ext,
    validate_image_content_type,
)
from common.utils.time_utils import get_beijing_now_naive

def _is_admin(user: User) -> bool:
    """判断用户是否为管理员"""
    return user.role == UserRole.ADMIN

router = APIRouter(prefix="/product-publish", tags=["商品发布"])


# ==================== Pydantic 请求 / 响应模型 ====================

class MaterialCreateRequest(BaseModel):
    """创建素材请求"""
    title: str = Field(..., min_length=1, max_length=200, description="商品标题")
    description: str = Field(..., min_length=1, description="商品描述")
    price: float = Field(..., gt=0, description="售价")
    original_price: Optional[float] = Field(None, description="原价（划线价）")
    category: Optional[str] = Field(None, max_length=100, description="商品分类")
    images: List[str] = Field(default=[], description="图片URL列表（最多9张）")
    delivery_method: str = Field("express", description="发货方式：express/pickup")
    postage: float = Field(0, ge=0, description="邮费，0表示包邮")
    address: Optional[str] = Field(None, max_length=200, description="宝贝所在地")
    brand: Optional[str] = Field(None, max_length=100, description="品牌")
    condition: str = Field("全新", description="成色")
    remark: Optional[str] = Field(None, max_length=500, description="备注（内部使用）")


class MaterialUpdateRequest(BaseModel):
    """更新素材请求（所有字段均可选）"""
    title: Optional[str] = Field(None, max_length=200)
    description: Optional[str] = None
    price: Optional[float] = Field(None, gt=0)
    original_price: Optional[float] = None
    category: Optional[str] = None
    images: Optional[List[str]] = None
    delivery_method: Optional[str] = None
    postage: Optional[float] = Field(None, ge=0)
    address: Optional[str] = None
    brand: Optional[str] = None
    condition: Optional[str] = None
    remark: Optional[str] = None


class PublishSingleRequest(BaseModel):
    """单品发布请求"""
    account_id: str = Field(..., description="闲鱼账号ID（cookie_id）")
    title: str = Field(..., min_length=1, max_length=200)
    description: str = Field(...)
    price: float = Field(..., gt=0)
    original_price: Optional[float] = None
    category: Optional[str] = Field(None, description="商品分类")
    images: List[str] = Field(..., min_length=1, description="图片本地路径列表（至少1张）")
    address: Optional[str] = None
    delivery_method: str = Field("express", description="发货方式：express/pickup")
    postage: float = Field(0, ge=0, description="邮费，0表示包邮")
    brand: Optional[str] = Field(None, description="品牌")
    condition: str = Field("全新", description="成色")


class BatchPublishRequest(BaseModel):
    """批量发布请求"""
    account_ids: List[str] = Field(..., min_length=1, description="账号ID列表")
    material_ids: List[int] = Field(..., min_length=1, description="素材ID列表")




# ==================== 1688 导入 / 登录态 ====================

class Save1688CookieRequest(BaseModel):
    """保存1688 Cookie请求"""
    cookie: str = Field(..., min_length=10, description="1688登录Cookie")


async def _ensure_1688_auth_table(session: AsyncSession) -> None:
    """确保1688登录态表存在"""
    await session.execute(text("""
        CREATE TABLE IF NOT EXISTS xy_1688_auth (
            id BIGINT NOT NULL AUTO_INCREMENT PRIMARY KEY,
            owner_id BIGINT NOT NULL,
            cookie LONGTEXT NULL,
            storage_state LONGTEXT NULL,
            status VARCHAR(20) NOT NULL DEFAULT 'unknown',
            last_check_at DATETIME NULL,
            created_at DATETIME NULL DEFAULT CURRENT_TIMESTAMP,
            updated_at DATETIME NULL DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP,
            UNIQUE KEY uk_1688_auth_owner (owner_id),
            KEY idx_1688_auth_owner (owner_id)
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci
    """))
    await session.commit()


@router.post("/1688/auth/cookie/save", response_model=ApiResponse)
async def save_1688_cookie(
    req: Save1688CookieRequest,
    current_user: User = Depends(get_current_user),
    session: AsyncSession = Depends(get_async_session),
):
    """保存1688登录Cookie"""
    cookie = (req.cookie or "").strip()
    if len(cookie) < 10:
        return ApiResponse(success=False, message="Cookie不能为空")

    await _ensure_1688_auth_table(session)

    await session.execute(
        text("""
            INSERT INTO xy_1688_auth (owner_id, cookie, status, last_check_at)
            VALUES (:owner_id, :cookie, 'saved', NOW())
            ON DUPLICATE KEY UPDATE
                cookie = VALUES(cookie),
                status = 'saved',
                last_check_at = NOW(),
                updated_at = NOW()
        """),
        {"owner_id": current_user.id, "cookie": cookie},
    )
    await session.commit()

    return ApiResponse(success=True, message="1688 Cookie已保存", data={"has_cookie": True, "status": "saved"})


@router.get("/1688/auth/status", response_model=ApiResponse)
async def get_1688_auth_status(
    current_user: User = Depends(get_current_user),
    session: AsyncSession = Depends(get_async_session),
):
    """检测是否已保存1688登录态"""
    await _ensure_1688_auth_table(session)

    row = (
        await session.execute(
            text("""
                SELECT cookie, status, last_check_at, updated_at
                FROM xy_1688_auth
                WHERE owner_id = :owner_id
                LIMIT 1
            """),
            {"owner_id": current_user.id},
        )
    ).mappings().first()

    has_cookie = bool(row and row.get("cookie"))
    return ApiResponse(
        success=True,
        message="查询成功",
        data={
            "has_cookie": has_cookie,
            "status": row.get("status") if row else "none",
            "last_check_at": str(row.get("last_check_at")) if row and row.get("last_check_at") else None,
            "updated_at": str(row.get("updated_at")) if row and row.get("updated_at") else None,
        },
    )


@router.post("/1688/auth/cookie/clear", response_model=ApiResponse)
async def clear_1688_cookie(
    current_user: User = Depends(get_current_user),
    session: AsyncSession = Depends(get_async_session),
):
    """清除1688登录Cookie"""
    await _ensure_1688_auth_table(session)

    await session.execute(
        text("""
            UPDATE xy_1688_auth
            SET cookie = NULL,
                storage_state = NULL,
                status = 'cleared',
                updated_at = NOW()
            WHERE owner_id = :owner_id
        """),
        {"owner_id": current_user.id},
    )
    await session.commit()

    return ApiResponse(success=True, message="1688 Cookie已清除", data={"has_cookie": False, "status": "cleared"})


# ==================== 素材库接口 ====================

@router.post("/materials", response_model=ApiResponse)
async def create_material(
    req: MaterialCreateRequest,
    current_user: User = Depends(get_current_active_user),
    session: AsyncSession = Depends(get_db_session),
) -> Dict[str, Any]:
    """创建商品素材"""
    svc = ProductMaterialService(session)
    material = await svc.create(current_user.id, req.model_dump())
    return ApiResponse(success=True, message="素材创建成功", data={"id": material.id})


@router.get("/materials", response_model=ApiResponse)
async def list_materials(
    page: int = Query(1, ge=1, description="页码"),
    page_size: int = Query(20, description="每页条数"),
    title: str = Query(None, description="标题模糊搜索"),
    category: str = Query(None, description="分类筛选"),
    condition: str = Query(None, description="成色筛选"),
    current_user: User = Depends(get_current_active_user),
    session: AsyncSession = Depends(get_db_session),
) -> Dict[str, Any]:
    """分页查询素材列表（管理员可查看所有用户的素材）"""
    svc = ProductMaterialService(session)
    # 管理员查看全部，普通用户只看自己的
    query_user_id = None if _is_admin(current_user) else current_user.id
    data = await svc.list_materials(
        query_user_id, page=page, page_size=page_size,
        title=title, category=category, condition=condition,
    )
    # 管理员场景：批量补充用户名
    if _is_admin(current_user) and data.get("list"):
        from sqlalchemy import select
        user_ids = list({m["user_id"] for m in data["list"]})
        stmt = select(User.id, User.username).where(User.id.in_(user_ids))
        rows = (await session.execute(stmt)).all()
        name_map = {r.id: r.username for r in rows}
        for m in data["list"]:
            m["username"] = name_map.get(m["user_id"], "未知用户")
    return ApiResponse(success=True, message="查询成功", data=data)


# ==================== 素材库导入 / 导出 ====================

# 素材库导出列。
# v1.2.2：取消中文“图片”列，避免导入时同时读取旧图片路径和 images/ 原图文件导致重复。
# 只有导出 ZIP 时才会额外写入英文 image_files 列，对应 ZIP 内 images/ 文件。
_MATERIAL_EXPORT_HEADERS = [
    "标题", "描述", "售价", "原价", "分类", "发货方式", "邮费", "宝贝所在地", "品牌", "成色", "备注",
]
_IMAGE_FILES_HEADER = "image_files"


def _split_images_cell(value: Any) -> List[str]:
    """解析 Excel 图片字段。支持换行、逗号、分号分隔。"""
    if value is None:
        return []
    text = str(value).strip()
    if not text:
        return []
    for sep in ["\r\n", "\n", "，", ",", ";", "；"]:
        text = text.replace(sep, "\n")
    images: List[str] = []
    for item in text.split("\n"):
        item = item.strip()
        if item:
            images.append(item)
    return images[:9]


def _to_float(value: Any, default: float | None = None) -> float | None:
    if value is None or str(value).strip() == "":
        return default
    try:
        return float(str(value).strip())
    except Exception:
        return default


def _safe_zip_name(value: Any, fallback: str = "image") -> str:
    """生成 ZIP 内安全文件名，避免中文/特殊符号导致解压异常。"""
    name = str(value or fallback).strip()
    name = unquote(name.split("?")[0].split("#")[0])
    name = os.path.basename(name) or fallback
    name = re.sub(r"[^0-9A-Za-z._\-\u4e00-\u9fa5]+", "_", name)
    return name[:120] or fallback


def _resolve_material_image_path(image_ref: str) -> Path | None:
    """尽量把素材图片引用解析为服务器本地文件路径。

    支持：
    1. 上传接口保存的绝对路径，例如 /app/data/uploads/products/xxx.jpg
    2. 前端预览 URL，例如 /static/uploads/products/xxx.jpg
    3. file:/// 开头的本地文件

    远程 http/https 图片不在服务器本地，导出 ZIP 时只保留原始链接，不主动下载。
    """
    if not image_ref:
        return None
    text = str(image_ref).strip()
    if not text:
        return None

    parsed = urlparse(text)
    if parsed.scheme in ("http", "https"):
        return None
    if parsed.scheme == "file":
        text = unquote(parsed.path)

    candidates: List[Path] = []
    if text.startswith("/static/uploads/products/"):
        try:
            from app.core.paths import get_upload_path
            candidates.append(Path(get_upload_path("products")) / os.path.basename(text))
        except Exception:
            pass
    candidates.append(Path(text))

    # 兼容部分环境把静态文件挂在项目 uploads 目录的情况。
    if "/static/uploads/products/" in text:
        try:
            from app.core.paths import get_upload_path
            candidates.append(Path(get_upload_path("products")) / os.path.basename(text))
        except Exception:
            pass

    for candidate in candidates:
        try:
            if candidate.exists() and candidate.is_file():
                return candidate
        except Exception:
            continue
    return None




async def _cleanup_unused_material_images(
    session: AsyncSession,
    image_refs: List[str] | None,
) -> Dict[str, Any]:
    """删除素材库中已不再被任何素材引用的本地商品图片文件。

    安全规则：
    1. 只处理本系统商品发布上传目录 products 下的本地图片。
    2. 远程 http/https 图片不删除。
    3. 如果图片仍被其他素材引用，不删除。
    4. 即使某张图片删除失败，也不影响素材删除/更新主流程。
    """
    refs = [str(x).strip() for x in (image_refs or []) if str(x or '').strip()]
    if not refs:
        return {"checked": 0, "deleted": 0, "skipped": 0, "failed": 0, "errors": []}

    try:
        from app.core.paths import get_upload_path
        products_dir = Path(get_upload_path("products")).resolve()
    except Exception as exc:
        logger.warning(f"素材图片清理跳过：无法获取 products 上传目录：{exc}")
        return {"checked": len(refs), "deleted": 0, "skipped": len(refs), "failed": 0, "errors": [str(exc)[:200]]}

    target_paths: Dict[Path, str] = {}
    for ref in refs:
        parsed = urlparse(ref)
        if parsed.scheme in ("http", "https"):
            continue
        local_path = _resolve_material_image_path(ref)
        if not local_path:
            continue
        try:
            resolved = local_path.resolve()
            # 只允许删除商品发布上传目录内的文件，防止误删系统文件。
            resolved.relative_to(products_dir)
            target_paths.setdefault(resolved, ref)
        except Exception:
            continue

    if not target_paths:
        return {"checked": len(refs), "deleted": 0, "skipped": len(refs), "failed": 0, "errors": []}

    remaining_paths: set[Path] = set()
    try:
        rows = (await session.execute(select(ProductMaterial.images))).scalars().all()
        for images in rows:
            for ref in (images or []):
                local_path = _resolve_material_image_path(str(ref))
                if not local_path:
                    continue
                try:
                    resolved = local_path.resolve()
                    resolved.relative_to(products_dir)
                    remaining_paths.add(resolved)
                except Exception:
                    continue
    except Exception as exc:
        logger.warning(f"素材图片清理跳过：查询剩余素材图片引用失败：{exc}")
        return {"checked": len(refs), "deleted": 0, "skipped": len(target_paths), "failed": 0, "errors": [str(exc)[:200]]}

    deleted = 0
    skipped = 0
    failed = 0
    errors: List[str] = []
    for path, ref in target_paths.items():
        if path in remaining_paths:
            skipped += 1
            continue
        try:
            if path.exists() and path.is_file():
                path.unlink()
                deleted += 1
                logger.info(f"素材图片已清理：{path}")
            else:
                skipped += 1
        except Exception as exc:
            failed += 1
            msg = f"{ref}: {exc}"
            errors.append(msg[:200])
            logger.warning(f"素材图片清理失败：{path}，原因：{exc}")

    return {"checked": len(refs), "deleted": deleted, "skipped": skipped, "failed": failed, "errors": errors[:20]}


def _append_material_rows(worksheet: Any, materials: List[Dict[str, Any]], image_export_map: Dict[int, List[str]] | None = None) -> None:
    """写入素材 Excel 行。

    v1.2.2：不再导出数据库里的“图片”路径列。
    ZIP 导出时只写英文 image_files，对应 ZIP 内 images/ 原图文件。
    """
    for m in materials:
        row = [
            m.get("title") or "",
            m.get("description") or "",
            m.get("price") or 0,
            m.get("original_price") or "",
            m.get("category") or "",
        ]
        if image_export_map is not None:
            row.append("\n".join(image_export_map.get(int(m.get("id") or 0), [])))
        row.extend([
            m.get("delivery_method") or "express",
            m.get("postage") or 0,
            m.get("address") or "",
            m.get("brand") or "",
            m.get("condition") or "全新",
            m.get("remark") or "",
        ])
        worksheet.append(row)


def _build_materials_excel(materials: List[Dict[str, Any]], image_export_map: Dict[int, List[str]] | None = None) -> io.BytesIO:
    """生成素材库 Excel。

    普通“仅导Excel”不包含图片列；
    ZIP 导出会增加英文 image_files 列，并且不再输出中文“图片/原图文件”列。
    """
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "商品素材库"
    headers = list(_MATERIAL_EXPORT_HEADERS)
    widths = [40, 60, 12, 12, 16, 14, 12, 30, 18, 14, 40]
    if image_export_map is not None:
        headers.insert(5, _IMAGE_FILES_HEADER)
        widths.insert(5, 80)
    worksheet.append(headers)
    _append_material_rows(worksheet, materials, image_export_map=image_export_map)
    for idx, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(idx)].width = width
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return output




def _normalize_import_image_key(value: Any) -> str:
    """规范化 Excel/文件夹上传中的图片相对路径，便于匹配 images/ 原图文件。"""
    text = str(value or "").strip()
    if not text:
        return ""
    text = unquote(text).replace("\\", "/")
    text = text.split("?")[0].split("#")[0].strip()
    while text.startswith("./"):
        text = text[2:]
    text = re.sub(r"/+", "/", text)
    return text.strip("/")


def _import_image_lookup_keys(value: Any) -> List[str]:
    """为一条图片引用生成多个候选匹配键。"""
    key = _normalize_import_image_key(value)
    if not key:
        return []
    keys = [key]
    basename = os.path.basename(key)
    if basename and basename not in keys:
        keys.append(basename)
    if "images/" in key:
        suffix = key[key.find("images/"):]
        if suffix and suffix not in keys:
            keys.append(suffix)
    return keys


def _hash_file_md5(path: Path) -> str | None:
    """计算本地图片 MD5。失败时返回 None。"""
    try:
        h = hashlib.md5()
        with path.open("rb") as f:
            for chunk in iter(lambda: f.read(1024 * 1024), b""):
                h.update(chunk)
        return h.hexdigest()
    except Exception:
        return None




async def save_uploaded_image(file: UploadFile, upload_dir: Path) -> tuple[Path, str, str]:
    """保存商品图片并按内容 MD5 去重。

    v1.2.3 修复：v1.2.2 调用了 save_uploaded_image 但没有定义，
    导致新建素材上传图片时报 name 'save_uploaded_image' is not defined。
    """
    upload_dir = Path(upload_dir)
    upload_dir.mkdir(parents=True, exist_ok=True)

    validate_image_content_type(file)
    content = await read_image_with_size_check(file)
    digest = hashlib.md5(content).hexdigest()
    ext = safe_image_ext(file.filename or "image.jpg")
    filename = f"{digest}{ext}"
    filepath = upload_dir / filename

    if filepath.exists():
        old_digest = _hash_file_md5(filepath)
        if old_digest == digest:
            return filepath, filename, digest
        filename = f"{digest}_{int(time.time())}{ext}"
        filepath = upload_dir / filename

    with filepath.open("wb") as f:
        f.write(content)

    return filepath, filename, digest


def _build_existing_product_image_hash_index(upload_dir: Path) -> Dict[str, str]:
    """扫描商品图片目录，建立 内容MD5 -> 静态URL 的索引。

    用于素材库导入去重：同一张原图之前已经存在时，直接复用旧 URL，
    不再生成新的 UUID 图片文件。
    """
    image_index: Dict[str, str] = {}
    if not upload_dir.exists():
        return image_index

    try:
        for path in upload_dir.iterdir():
            if not path.is_file():
                continue
            if path.suffix.lower() not in {".jpg", ".jpeg", ".png", ".gif", ".webp", ".bmp"}:
                continue
            digest = _hash_file_md5(path)
            if digest:
                image_index.setdefault(digest, f"/static/uploads/products/{path.name}")
    except Exception as exc:
        logger.warning(f"素材导入图片去重索引建立失败，将仅做本批次去重：{exc}")
    return image_index


async def _save_import_image_files(image_files: Optional[List[UploadFile]]) -> Dict[str, str]:
    """保存导入时一起上传的 images 文件夹图片，并返回 原相对路径/文件名 -> 新系统图片 URL。

    v1.2.1 修复：导入图片按文件内容 MD5 去重。
    - 服务器 products 目录已经存在同图：复用旧 URL，不重复保存。
    - 本次 images 文件夹里同图出现多次：复用第一次保存/匹配结果。
    - 新图：用 MD5 文件名保存，后续再次导入同一张图会稳定复用。
    """
    image_map: Dict[str, str] = {}
    if not image_files:
        return image_map

    from app.core.paths import get_upload_path
    upload_dir = Path(get_upload_path("products"))
    upload_dir.mkdir(parents=True, exist_ok=True)

    existing_image_index = _build_existing_product_image_hash_index(upload_dir)
    batch_image_index: Dict[str, str] = {}
    saved_count = 0
    reused_count = 0

    for image in image_files:
        if not image or not image.filename:
            continue
        original_name = image.filename
        try:
            validate_image_content_type(image)
            content = await read_image_with_size_check(image)
            digest = hashlib.md5(content).hexdigest()

            saved_url = existing_image_index.get(digest) or batch_image_index.get(digest)
            if saved_url:
                reused_count += 1
            else:
                ext = safe_image_ext(original_name)
                filename = f"{digest}{ext}"
                filepath = upload_dir / filename

                # 极小概率同名存在但未进入索引；如果内容相同复用，否则加时间后缀避免覆盖。
                if filepath.exists():
                    old_digest = _hash_file_md5(filepath)
                    if old_digest == digest:
                        saved_url = f"/static/uploads/products/{filename}"
                        reused_count += 1
                    else:
                        filename = f"{digest}_{int(time.time())}{ext}"
                        filepath = upload_dir / filename

                if not saved_url:
                    with filepath.open("wb") as f:
                        f.write(content)
                    saved_url = f"/static/uploads/products/{filename}"
                    saved_count += 1

                existing_image_index.setdefault(digest, saved_url)
                batch_image_index.setdefault(digest, saved_url)

            for key in _import_image_lookup_keys(original_name):
                image_map.setdefault(key, saved_url)
        except ImageUploadError as exc:
            logger.warning(f"素材导入图片跳过：{original_name}，原因：{exc.message}")
        except Exception as exc:
            logger.warning(f"素材导入图片保存失败：{original_name}，原因：{exc}")

    logger.info(f"素材导入图片处理完成：新保存 {saved_count} 张，复用 {reused_count} 张，匹配键 {len(image_map)} 个")
    return image_map


def _resolve_imported_images(
    image_files_cell: Any,
    image_map: Dict[str, str],
) -> List[str]:
    """根据 Excel 的 image_files 列 + 上传的 images 文件夹，生成素材图片 URL。

    v1.2.2：只读取英文 image_files 列。
    不再读取“图片”列，也不再保留 Excel 里的旧 /static/uploads/products 路径，
    避免重复导入时旧路径和新上传图片互相叠加。
    """
    refs = _split_images_cell(image_files_cell)
    result: List[str] = []
    for ref in refs:
        matched = None
        for key in _import_image_lookup_keys(ref):
            matched = image_map.get(key)
            if matched:
                break
        if matched and matched not in result:
            result.append(matched)
        if len(result) >= 9:
            break
    return result[:9]


@router.get("/materials/export")
async def export_materials(
    title: str = Query(None, description="标题模糊搜索"),
    category: str = Query(None, description="分类筛选"),
    condition: str = Query(None, description="成色筛选"),
    current_user: User = Depends(get_current_active_user),
    session: AsyncSession = Depends(get_db_session),
):
    """导出商品素材库为 Excel。管理员导出全部，普通用户只导出自己的素材。"""
    svc = ProductMaterialService(session)
    query_user_id = None if _is_admin(current_user) else current_user.id
    data = await svc.list_materials(
        query_user_id,
        page=1,
        page_size=1000,
        title=title,
        category=category,
        condition=condition,
    )

    output = _build_materials_excel(data.get("list", []))
    filename = f"product_materials_{int(time.time())}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/materials/export-with-images")
async def export_materials_with_images(
    title: str = Query(None, description="标题模糊搜索"),
    category: str = Query(None, description="分类筛选"),
    condition: str = Query(None, description="成色筛选"),
    current_user: User = Depends(get_current_active_user),
    session: AsyncSession = Depends(get_db_session),
):
    """导出商品素材库 ZIP：materials.xlsx + images/ 原图。

    v1.2.2：Excel 不再包含中文“图片/原图文件”列；只包含英文 image_files 列，
    对应 ZIP 内 images/ 下的文件。远程 URL 无法保证服务器能拿到原文件，会写入 missing_images.txt。
    """
    svc = ProductMaterialService(session)
    query_user_id = None if _is_admin(current_user) else current_user.id
    data = await svc.list_materials(
        query_user_id,
        page=1,
        page_size=1000,
        title=title,
        category=category,
        condition=condition,
    )
    materials = data.get("list", [])

    image_export_map: Dict[int, List[str]] = {}
    missing_lines: List[str] = []
    zip_buffer = io.BytesIO()

    with zipfile.ZipFile(zip_buffer, mode="w", compression=zipfile.ZIP_DEFLATED) as zf:
        for m in materials:
            material_id = int(m.get("id") or 0)
            title_text = _safe_zip_name(m.get("title") or f"material_{material_id}", f"material_{material_id}")
            image_export_map.setdefault(material_id, [])
            for idx, image_ref in enumerate((m.get("images") or [])[:9], start=1):
                local_path = _resolve_material_image_path(str(image_ref))
                if not local_path:
                    missing_lines.append(f"素材ID {material_id}《{m.get('title') or ''}》第{idx}张：未找到本地原图，原始引用：{image_ref}")
                    continue
                suffix = local_path.suffix or ".jpg"
                image_name = f"{material_id}_{idx}_{title_text}{suffix}"
                image_name = _safe_zip_name(image_name, f"{material_id}_{idx}{suffix}")
                arcname = f"images/{material_id}/{image_name}"
                try:
                    zf.write(local_path, arcname)
                    image_export_map[material_id].append(arcname)
                except Exception as exc:
                    missing_lines.append(f"素材ID {material_id}《{m.get('title') or ''}》第{idx}张：写入ZIP失败：{exc}，原始引用：{image_ref}")

        excel_buffer = _build_materials_excel(materials, image_export_map=image_export_map)
        zf.writestr("materials.xlsx", excel_buffer.getvalue())
        readme = (
            "商品素材库原图导出说明\n"
            "1. materials.xlsx 是素材数据表。\n"
            "2. images/ 目录内是能在服务器本地找到的原图。\n"
            "3. Excel 不再导出中文“图片/原图文件”列，只保留英文 image_files 列。\n"
            "4. image_files 列对应本 ZIP 内 images/ 目录下的图片文件。\n"
            "5. 如果某张图片是远程 URL、已被删除、或服务器本地不存在，会记录在 missing_images.txt。\n"
        )
        zf.writestr("README.txt", readme)
        if missing_lines:
            zf.writestr("missing_images.txt", "\n".join(missing_lines))

    zip_buffer.seek(0)
    filename = f"product_materials_with_images_{int(time.time())}.zip"
    return StreamingResponse(
        zip_buffer,
        media_type="application/zip",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.post("/materials/import", response_model=ApiResponse)
async def import_materials(
    file: UploadFile = File(...),
    image_files: Optional[List[UploadFile]] = File(None),
    current_user: User = Depends(get_current_active_user),
    session: AsyncSession = Depends(get_db_session),
) -> Dict[str, Any]:
    """从 Excel 导入商品素材。

    支持两种导入：
    1. 只上传 Excel：只导入文字/价格等信息，不导入图片。
    2. 上传 Excel + images 文件夹：只按英文 image_files 列匹配 images/ 原图并保存/复用到本系统。
    必填列：标题、描述、售价。导入为追加方式，不覆盖原素材。
    """
    if not file.filename or not file.filename.lower().endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=400, detail="请上传 Excel 文件（.xlsx 或 .xls）")

    contents = await file.read()
    try:
        workbook = load_workbook(io.BytesIO(contents), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Excel 文件读取失败：{exc}") from exc

    worksheet = workbook.active
    rows = list(worksheet.iter_rows(values_only=True))
    if not rows:
        return ApiResponse(success=False, message="Excel 文件为空", data={"created": 0, "failed": 0})

    header = [str(cell).strip() if cell is not None else "" for cell in rows[0]]
    aliases = {
        "标题": ["标题", "商品标题", "title"],
        "描述": ["描述", "商品描述", "description"],
        "售价": ["售价", "价格", "price"],
        "原价": ["原价", "划线价", "original_price"],
        "分类": ["分类", "商品分类", "category"],
        # v1.2.2：导入只认英文 image_files。
        # 不再读取中文“图片”列，避免旧 /static 路径和 images 文件夹重复叠加。
        "image_files": ["image_files", "original_image_files", "original_images"],
        "发货方式": ["发货方式", "delivery_method"],
        "邮费": ["邮费", "postage"],
        "宝贝所在地": ["宝贝所在地", "所在地", "地址", "address"],
        "品牌": ["品牌", "brand"],
        "成色": ["成色", "condition"],
        "备注": ["备注", "remark"],
    }

    def col(name: str) -> int | None:
        for alias in aliases[name]:
            if alias in header:
                return header.index(alias)
        return None

    required = ["标题", "描述", "售价"]
    missing = [name for name in required if col(name) is None]
    if missing:
        return ApiResponse(success=False, message=f"Excel 缺少必要列：{', '.join(missing)}", data={"created": 0, "failed": 0})

    created = 0
    failed = 0
    errors: List[Dict[str, Any]] = []
    svc = ProductMaterialService(session)
    image_map = await _save_import_image_files(image_files)

    for row_no, row in enumerate(rows[1:5001], start=2):
        try:
            def cell(name: str) -> Any:
                idx = col(name)
                return row[idx] if idx is not None and len(row) > idx else None

            title_value = str(cell("标题") or "").strip()
            description_value = str(cell("描述") or "").strip()
            price_value = _to_float(cell("售价"), None)
            if not title_value and not description_value and price_value is None:
                continue
            if not title_value or not description_value or price_value is None or price_value <= 0:
                failed += 1
                errors.append({"row": row_no, "reason": "标题、描述、售价为必填，且售价必须大于0"})
                continue

            original_price = _to_float(cell("原价"), None)
            postage = _to_float(cell("邮费"), 0) or 0
            delivery_method = str(cell("发货方式") or "express").strip() or "express"
            if delivery_method not in ("express", "pickup"):
                delivery_method = "express"

            await svc.create(current_user.id, {
                "title": title_value[:200],
                "description": description_value,
                "price": price_value,
                "original_price": original_price,
                "category": str(cell("分类") or "").strip()[:100] or None,
                "images": _resolve_imported_images(cell("image_files"), image_map),
                "delivery_method": delivery_method,
                "postage": postage,
                "address": str(cell("宝贝所在地") or "").strip()[:200] or None,
                "brand": str(cell("品牌") or "").strip()[:100] or None,
                "condition": str(cell("成色") or "全新").strip()[:20] or "全新",
                "remark": str(cell("备注") or "").strip()[:500] or None,
            })
            created += 1
        except Exception as exc:
            failed += 1
            errors.append({"row": row_no, "reason": str(exc)[:200]})

    matched_image_count = len(set(image_map.values()))
    message = f"导入完成：成功 {created} 条"
    if image_map:
        message += f"，已匹配导入图片 {matched_image_count} 张（同图自动复用，不重复保存）"
    if failed:
        message += f"，失败 {failed} 条"
    return ApiResponse(
        success=True,
        message=message,
        data={"created": created, "failed": failed, "matched_images": matched_image_count, "errors": errors[:20]},
    )


class BatchDeleteRequest(BaseModel):
    """批量删除素材请求"""
    ids: List[int] = Field(..., min_length=1, description="素材ID列表")


@router.post("/materials/batch-delete", response_model=ApiResponse)
async def batch_delete_materials(
    req: BatchDeleteRequest,
    current_user: User = Depends(get_current_active_user),
    session: AsyncSession = Depends(get_db_session),
) -> Dict[str, Any]:
    """批量删除素材（管理员可删除任意素材），并清理不再被引用的本地图片。"""
    svc = ProductMaterialService(session)
    query_user_id = None if _is_admin(current_user) else current_user.id

    before_images: List[str] = []
    try:
        conds = [ProductMaterial.id.in_(req.ids)]
        if query_user_id is not None:
            conds.append(ProductMaterial.user_id == query_user_id)
        rows = (await session.execute(select(ProductMaterial.images).where(*conds))).scalars().all()
        for images in rows:
            before_images.extend(images or [])
    except Exception as exc:
        logger.warning(f"批量删除素材前获取图片引用失败，将只删除素材记录：{exc}")

    count = await svc.batch_delete(req.ids, query_user_id)
    cleanup = await _cleanup_unused_material_images(session, before_images) if count else {"deleted": 0, "failed": 0}

    message = f"成功删除 {count} 条素材"
    if cleanup.get("deleted"):
        message += f"，已清理 {cleanup['deleted']} 张本地图片"
    if cleanup.get("failed"):
        message += f"，{cleanup['failed']} 张图片清理失败"
    return ApiResponse(success=True, message=message, data={"deleted_count": count, "image_cleanup": cleanup})


@router.get("/materials/{material_id}", response_model=ApiResponse)
async def get_material(
    material_id: int,
    current_user: User = Depends(get_current_active_user),
    session: AsyncSession = Depends(get_db_session),
) -> Dict[str, Any]:
    """获取单条素材详情（管理员可访问任意素材）"""
    svc = ProductMaterialService(session)
    query_user_id = None if _is_admin(current_user) else current_user.id
    material = await svc.get(material_id, query_user_id)
    if not material:
        return ApiResponse(success=False, message="素材不存在或无权访问")
    from app.services.product_publish_service import _material_to_dict
    return ApiResponse(success=True, message="查询成功", data=_material_to_dict(material))


@router.put("/materials/{material_id}", response_model=ApiResponse)
async def update_material(
    material_id: int,
    req: MaterialUpdateRequest,
    current_user: User = Depends(get_current_active_user),
    session: AsyncSession = Depends(get_db_session),
) -> Dict[str, Any]:
    """更新素材信息（管理员可修改任意素材），图片被替换时清理不再被引用的旧本地图片。"""
    svc = ProductMaterialService(session)
    query_user_id = None if _is_admin(current_user) else current_user.id
    old_material = await svc.get(material_id, query_user_id)
    if not old_material:
        return ApiResponse(success=False, message="素材不存在或无权修改")

    old_images = list(old_material.images or [])
    update_data = {k: v for k, v in req.model_dump().items() if v is not None}
    updated = await svc.update(material_id, query_user_id, update_data)
    if not updated:
        return ApiResponse(success=False, message="素材不存在或无权修改")

    cleanup = {"deleted": 0, "failed": 0}
    if "images" in update_data:
        new_images = set(str(x) for x in (update_data.get("images") or []))
        removed_images = [img for img in old_images if str(img) not in new_images]
        cleanup = await _cleanup_unused_material_images(session, removed_images)

    message = "素材更新成功"
    if cleanup.get("deleted"):
        message += f"，已清理 {cleanup['deleted']} 张旧图片"
    if cleanup.get("failed"):
        message += f"，{cleanup['failed']} 张旧图片清理失败"
    return ApiResponse(success=True, message=message, data={"image_cleanup": cleanup})


@router.delete("/materials/{material_id}", response_model=ApiResponse)
async def delete_material(
    material_id: int,
    current_user: User = Depends(get_current_active_user),
    session: AsyncSession = Depends(get_db_session),
) -> Dict[str, Any]:
    """删除素材（管理员可删除任意素材），并清理不再被引用的本地图片。"""
    svc = ProductMaterialService(session)
    query_user_id = None if _is_admin(current_user) else current_user.id
    material = await svc.get(material_id, query_user_id)
    if not material:
        return ApiResponse(success=False, message="素材不存在或无权删除")
    before_images = list(material.images or [])

    deleted = await svc.delete(material_id, query_user_id)
    if not deleted:
        return ApiResponse(success=False, message="素材不存在或无权删除")

    cleanup = await _cleanup_unused_material_images(session, before_images)
    message = "素材删除成功"
    if cleanup.get("deleted"):
        message += f"，已清理 {cleanup['deleted']} 张本地图片"
    if cleanup.get("failed"):
        message += f"，{cleanup['failed']} 张图片清理失败"
    return ApiResponse(success=True, message=message, data={"image_cleanup": cleanup})


# ==================== 发布接口 ====================

@router.post("/publish/single", response_model=ApiResponse)
async def publish_single(
    req: PublishSingleRequest,
    current_user: User = Depends(get_current_active_user),
    session: AsyncSession = Depends(get_db_session),
) -> Dict[str, Any]:
    """单品发布（同步执行，等待 Playwright 完成后返回结果）
    
    注意：发布操作会启动无头浏览器，耗时约 30-60 秒，请前端设置合适的超时时间。
    """
    svc = PublishExecutorService(session)
    result = await svc.publish_single(
        user_id=current_user.id,
        account_id=req.account_id,
        item_data=req.model_dump(),
    )
    return ApiResponse(
        success=result.get("success", False),
        message=result.get("message", ""),
        data={
            "item_url": result.get("item_url"),
            "item_id": result.get("item_id"),
            "log_id": result.get("log_id"),
            "sync_status": result.get("sync_status"),
            "sync_message": result.get("sync_message"),
            "sync_total_count": result.get("sync_total_count"),
            "sync_saved_count": result.get("sync_saved_count"),
        },
    )


@router.post("/publish/batch", response_model=ApiResponse)
async def publish_batch(
    req: BatchPublishRequest,
    background_tasks: BackgroundTasks,
    current_user: User = Depends(get_current_active_user),
    session: AsyncSession = Depends(get_db_session),
) -> Dict[str, Any]:
    """批量发布（后台异步执行，立即返回 batch_id）
    
    前端通过 GET /publish/batch/{batch_id}/status 查询进度。
    后台会按账号循环，每个账号依次发布所有素材，复用同一浏览器实例。
    """
    mat_svc = ProductMaterialService(session)
    from app.services.product_publish_service import _material_to_dict
    materials = [_material_to_dict(m) for m in await mat_svc.list_by_ids(req.material_ids, current_user.id)]

    if not materials:
        return ApiResponse(success=False, message="没有找到有效的素材")

    batch_id = str(uuid.uuid4())
    await PublishBatchStatusService.init_batch(
        batch_id=batch_id,
        account_ids=req.account_ids,
        material_count=len(materials),
        user_id=current_user.id,
    )

    # 创建后台任务
    background_tasks.add_task(
        _run_batch_publish_background,
        user_id=current_user.id,
        account_ids=req.account_ids,
        materials=materials,
        batch_id=batch_id,
    )

    return ApiResponse(
        success=True,
        message=f"批量发布任务已提交，共 {len(req.account_ids)} 个账号 × {len(materials)} 件商品",
        data={
            "batch_id": batch_id,
            "total": len(req.account_ids) * len(materials),
        },
    )


@router.post("/publish/batch/{batch_id}/cancel", response_model=ApiResponse)
async def cancel_batch_publish(
    batch_id: str,
    current_user: User = Depends(get_current_active_user),
    session: AsyncSession = Depends(get_db_session),
) -> Dict[str, Any]:
    """请求停止正在执行的批量发布任务。

    说明：
    - 等待中的队列任务由前端直接取消，不会进入这里。
    - 已提交到后端的任务会被标记为取消；后台发布逻辑会在账号/商品之间以及单品发布执行中轮询该标记并停止。
    - 已经发布成功的商品不会回滚；尚未开始的商品不会继续发布。
    """
    from common.models.publish_log import PublishLog

    snapshot = await PublishBatchStatusService.get_batch_snapshot(batch_id)
    if snapshot is not None and snapshot.get("user_id") not in (None, current_user.id):
        return ApiResponse(success=False, message="无权停止该批量发布任务")

    cancelled = await PublishBatchStatusService.request_cancel(batch_id, "批量发布已手动停止")

    # 尽量把当前仍处于 publishing/pending 的日志标记为失败，方便发布日志里看到停止原因。
    # 后台如果刚好完成某条发布，可能会覆盖成 success；这是正常竞争，以实际发布结果为准。
    try:
        stmt = (
            update(PublishLog)
            .where(
                PublishLog.batch_id == batch_id,
                PublishLog.user_id == current_user.id,
                PublishLog.status.in_(["pending", "publishing"]),
            )
            .values(status="failed", error_message="批量发布已手动停止")
        )
        await session.execute(stmt)
        await session.commit()
    except Exception:
        await session.rollback()

    if not cancelled:
        return ApiResponse(success=False, message="任务不存在、已完成或状态已失效")
    return ApiResponse(success=True, message="已请求停止批量发布任务", data={"batch_id": batch_id})


@router.get("/publish/batch/{batch_id}/status", response_model=ApiResponse)
async def get_batch_status(
    batch_id: str,
    current_user: User = Depends(get_current_active_user),
    session: AsyncSession = Depends(get_db_session),
) -> Dict[str, Any]:
    """查询批量发布任务进度"""
    from sqlalchemy import select, func
    from common.models.publish_log import PublishLog

    unknown_sync_message = "批量任务同步状态缓存不存在，无法判断自动获取商品结果"

    stmt = select(
        PublishLog.status,
        func.count().label("cnt"),
    ).where(
        PublishLog.batch_id == batch_id,
        PublishLog.user_id == current_user.id,
    ).group_by(PublishLog.status)

    rows = (await session.execute(stmt)).all()
    counts = {r.status: r.cnt for r in rows}

    account_stmt = select(
        PublishLog.account_id,
        PublishLog.status,
        func.count().label("cnt"),
    ).where(
        PublishLog.batch_id == batch_id,
        PublishLog.user_id == current_user.id,
    ).group_by(PublishLog.account_id, PublishLog.status)
    account_rows = (await session.execute(account_stmt)).all()

    account_count_map: Dict[str, Dict[str, int]] = {}
    for row in account_rows:
        status_map = account_count_map.setdefault(row.account_id, {})
        status_map[row.status] = int(row.cnt)

    total = sum(counts.values())
    success = counts.get("success", 0)
    failed = counts.get("failed", 0)
    publishing = counts.get("publishing", 0)
    pending = counts.get("pending", 0)
    batch_snapshot = await PublishBatchStatusService.get_batch_snapshot(batch_id)

    if batch_snapshot is None:
        if total == 0:
            return ApiResponse(success=False, message="批量任务不存在或状态已失效")
        return ApiResponse(success=False, message="批量任务状态已失效，请到发布日志查看执行结果")

    account_statuses: List[Dict[str, Any]] = []
    batch_cancelled = bool(batch_snapshot.get("cancelled")) if batch_snapshot else False
    batch_cancel_message = batch_snapshot.get("cancel_message") if batch_snapshot else None

    if batch_snapshot:
        material_count = int(batch_snapshot.get("material_count") or 0)
        account_order = batch_snapshot.get("account_order") or []
        account_sync_map = batch_snapshot.get("accounts") or {}
        expected_total = material_count * len(account_order)
        if expected_total > total:
            total = expected_total
            pending = max(total - success - failed - publishing, 0)

        for account_id in account_order:
            status_map = account_count_map.get(account_id, {})
            account_total = material_count if material_count > 0 else sum(status_map.values())
            account_success = int(status_map.get("success", 0))
            account_failed = int(status_map.get("failed", 0))
            account_publishing = int(status_map.get("publishing", 0))
            account_pending = max(account_total - account_success - account_failed - account_publishing, 0)
            sync_info = account_sync_map.get(account_id, {})
            account_statuses.append(
                {
                    "account_id": account_id,
                    "total": account_total,
                    "success": account_success,
                    "failed": account_failed,
                    "publishing": account_publishing,
                    "pending": account_pending,
                    "sync_status": sync_info.get("sync_status", "pending"),
                    "sync_message": sync_info.get("sync_message", "等待该账号发布完成后自动获取商品"),
                    "sync_total_count": int(sync_info.get("sync_total_count") or 0),
                    "sync_saved_count": int(sync_info.get("sync_saved_count") or 0),
                }
            )

        extra_account_ids = [account_id for account_id in account_count_map.keys() if account_id not in set(account_order)]
        for account_id in extra_account_ids:
            status_map = account_count_map.get(account_id, {})
            account_total = sum(status_map.values())
            account_success = int(status_map.get("success", 0))
            account_failed = int(status_map.get("failed", 0))
            account_publishing = int(status_map.get("publishing", 0))
            account_pending = int(status_map.get("pending", 0))
            account_statuses.append(
                {
                    "account_id": account_id,
                    "total": account_total,
                    "success": account_success,
                    "failed": account_failed,
                    "publishing": account_publishing,
                    "pending": account_pending,
                    "sync_status": "unknown",
                    "sync_message": unknown_sync_message,
                    "sync_total_count": 0,
                    "sync_saved_count": 0,
                }
            )
    sync_finished = all(
        account_status.get("sync_status") in {"success", "failed", "skipped", "unknown"}
        for account_status in account_statuses
    ) if account_statuses else True

    if batch_cancelled:
        publishing = 0
        pending = 0

    return ApiResponse(
        success=True,
        message="查询成功",
        data={
            "batch_id": batch_id,
            "total": total,
            "success": success,
            "failed": failed,
            "publishing": publishing,
            "pending": pending,
            "finished": (batch_cancelled or (total > 0 and (publishing + pending) == 0 and sync_finished)),
            "cancelled": batch_cancelled,
            "cancel_message": batch_cancel_message,
            "account_statuses": account_statuses,
        },
    )


# ==================== 发布日志接口 ====================

@router.get("/logs", response_model=ApiResponse)
async def list_publish_logs(
    page: int = Query(1, ge=1),
    page_size: int = Query(20),
    account_id: Optional[str] = Query(None, description="按账号过滤"),
    status: Optional[str] = Query(None, description="按状态过滤：pending/publishing/success/failed"),
    current_user: User = Depends(get_current_active_user),
    session: AsyncSession = Depends(get_db_session),
) -> Dict[str, Any]:
    """分页查询发布日志（管理员可查看所有用户的发布日志）"""
    svc = PublishLogService(session)
    # 管理员查看全部，普通用户只看自己的
    query_user_id = None if _is_admin(current_user) else current_user.id
    data = await svc.list_logs(
        user_id=query_user_id,
        page=page,
        page_size=page_size,
        account_id=account_id,
        status=status,
    )
    # 管理员场景：批量补充用户名
    if _is_admin(current_user) and data.get("list"):
        from sqlalchemy import select
        user_ids = list({log["user_id"] for log in data["list"]})
        stmt = select(User.id, User.username).where(User.id.in_(user_ids))
        rows = (await session.execute(stmt)).all()
        name_map = {r.id: r.username for r in rows}
        for log in data["list"]:
            log["username"] = name_map.get(log["user_id"], "未知用户")
    return ApiResponse(success=True, message="查询成功", data=data)


@router.delete("/logs/clear", response_model=ApiResponse)
async def clear_publish_logs(
    current_user: User = Depends(get_current_active_user),
    session: AsyncSession = Depends(get_db_session),
) -> Dict[str, Any]:
    """清空发布日志（只清空30天前的数据）"""
    from datetime import timedelta

    from loguru import logger
    from sqlalchemy import delete

    from common.models.publish_log import PublishLog

    try:
        thirty_days_ago = get_beijing_now_naive() - timedelta(days=30)
        stmt = delete(PublishLog).where(
            PublishLog.user_id == current_user.id,
            PublishLog.created_at < thirty_days_ago,
        )

        result = await session.execute(stmt)
        await session.commit()

        deleted_count = result.rowcount or 0
        logger.info(f"[发布日志] 用户 {current_user.id} 已清空 {deleted_count} 条30天前的日志")
        return ApiResponse(
            success=True,
            message=f"已清空 {deleted_count} 条30天前的发布日志",
        )
    except Exception as e:
        await session.rollback()
        logger.error(f"[发布日志] 清空日志失败: {e}")
        return ApiResponse(success=False, message=f"清空发布日志失败: {str(e)}")


# ==================== 图片上传接口 ====================

@router.post("/upload/images", response_model=ApiResponse)
async def upload_product_images(
    files: List[UploadFile] = File(...),
    current_user: User = Depends(get_current_active_user),
) -> Dict[str, Any]:
    """上传商品图片（支持多张，最多9张，每张最大5MB）

    返回本地文件路径列表，这些路径将直接传给 Playwright 的 set_input_files。
    """
    from app.core.paths import get_upload_path

    upload_dir = get_upload_path("products")

    if len(files) > 9:
        return ApiResponse(success=False, message="最多上传9张图片")

    saved_paths: List[str] = []
    saved_urls: List[str] = []

    for file in files:
        try:
            filepath, filename, _ = await save_uploaded_image(
                file,
                upload_dir,
            )
        except ImageUploadError as exc:
            # 在消息里带上具体哪张图片出错，方便前端展示
            return ApiResponse(
                success=False,
                message=f"文件 {file.filename}: {exc.message}",
            )

        saved_paths.append(str(filepath))                          # 绝对路径，用于 Playwright
        saved_urls.append(f"/static/uploads/products/{filename}")  # URL，用于前端预览

    return ApiResponse(
        success=True,
        message=f"成功上传 {len(saved_paths)} 张图片",
        data={"paths": saved_paths, "urls": saved_urls},
    )


# ==================== 后台任务函数 ====================

async def _run_batch_publish_background(
    user_id: int,
    account_ids: List[str],
    materials: List[dict],
    batch_id: str,
) -> None:
    """后台异步执行批量发布任务"""
    from common.db.session import async_session_maker
    from loguru import logger
    import traceback

    async with async_session_maker() as session:
        svc = PublishExecutorService(session)
        try:
            # 直接将 batch_id 传给 service，确保日志与路由返回值一致
            await svc.batch_publish(
                user_id=user_id,
                account_ids=account_ids,
                materials=materials,
                batch_id=batch_id,
            )
        except Exception as e:
            logger.error(f"批量发布后台任务异常: {e}\n{traceback.format_exc()}")
            await PublishBatchStatusService.clear_batch(batch_id)